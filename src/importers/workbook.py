"""Stage 1: read workbook metadata and raw cell values.

Synchronous (openpyxl); callers run it via `asyncio.to_thread`.
"""

import hashlib
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class RawRow:
    row_number: int  # 1-based Excel row number, for error messages
    values: dict[str, object]


@dataclass
class RawSheet:
    name: str
    headers: list[str]
    rows: list[RawRow] = field(default_factory=list)
    blank_rows_skipped: int = 0


@dataclass
class RawWorkbook:
    path: str
    sha256: str
    sheet_names: list[str]
    sheets: dict[str, RawSheet]


def _header(value: object) -> str:
    return "" if value is None else str(value).strip()


def read_workbook(path: Path) -> RawWorkbook:
    """Read every sheet as header + row dicts. No interpretation of values happens here."""
    data = path.read_bytes()
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheets: dict[str, RawSheet] = {}
        for ws in workbook.worksheets:
            rows = ws.iter_rows(values_only=True)
            first = next(rows, None)
            headers = [_header(v) for v in first] if first else []
            sheet = RawSheet(name=ws.title, headers=headers)
            for offset, values in enumerate(rows, start=2):
                if all(v is None or (isinstance(v, str) and not v.strip()) for v in values):
                    sheet.blank_rows_skipped += 1
                    continue
                sheet.rows.append(RawRow(offset, dict(zip(headers, values, strict=False))))
            sheets[ws.title] = sheet
        return RawWorkbook(
            path=str(path),
            sha256=hashlib.sha256(data).hexdigest(),
            sheet_names=list(sheets),
            sheets=sheets,
        )
    finally:
        workbook.close()
